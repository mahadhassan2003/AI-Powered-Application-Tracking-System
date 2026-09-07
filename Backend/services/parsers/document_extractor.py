import logging
from pathlib import Path
from .exceptions import ResumeParsingError

try:
    import PyPDF2
except ImportError:
    PyPDF2 = None

try:
    from docx import Document
except ImportError:
    Document = None

logger = logging.getLogger(__name__)

class DocumentExtractor:
    @staticmethod
    def extract_text_from_bytes(file_content: bytes, filename: str) -> str:
        import tempfile
        
        file_extension = Path(filename).suffix.lower()
        with tempfile.NamedTemporaryFile(suffix=file_extension, delete=False) as tmp_file:
            tmp_file.write(file_content)
            tmp_path = tmp_file.name
        
        try:
            if file_extension == '.pdf':
                text = DocumentExtractor.extract_text_from_pdf(tmp_path)
            elif file_extension == '.docx':
                text = DocumentExtractor.extract_text_from_docx(tmp_path)
            else:
                raise ResumeParsingError(f"Unsupported file type: {file_extension}")
            return text
        finally:
            Path(tmp_path).unlink(missing_ok=True)

    @staticmethod
    def extract_text_from_pdf(file_path: str) -> str:
        try:
            if PyPDF2 is None:
                raise ResumeParsingError("PyPDF2 library not installed.")
                
            text = ""
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page_num in range(len(pdf_reader.pages)):
                    page = pdf_reader.pages[page_num]
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
            
            if not text.strip():
                raise ResumeParsingError("No text could be extracted from PDF")
            return text.strip()
        except Exception as e:
            logger.error(f"PDF text extraction failed for {file_path}: {e}")
            raise ResumeParsingError(f"Failed to extract text from PDF: {str(e)}")

    @staticmethod
    def extract_text_from_docx(file_path: str) -> str:
        try:
            if Document is None:
                raise ResumeParsingError("python-docx library not installed.")
                
            doc = Document(file_path)
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
                
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        text += cell.text + "\n"
            
            if not text.strip():
                raise ResumeParsingError("No text could be extracted from DOCX")
            return text.strip()
        except Exception as e:
            logger.error(f"DOCX text extraction failed for {file_path}: {e}")
            raise ResumeParsingError(f"Failed to extract text from DOCX: {str(e)}")

    @staticmethod
    def extract_text_from_file(file_path: str) -> str:
        file_path_obj = Path(file_path)
        file_extension = file_path_obj.suffix.lower()
        
        if file_extension == '.pdf':
            return DocumentExtractor.extract_text_from_pdf(str(file_path))
        elif file_extension == '.docx':
            return DocumentExtractor.extract_text_from_docx(str(file_path))
        else:
            raise ResumeParsingError(f"Unsupported file type: {file_extension}")
